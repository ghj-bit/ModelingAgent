"""Collect one co-evolution run's rubric, dialogue and judge evidence into one workbook.

The co-evolution arm evolves two things at once -- the interaction policy and the
rubric the critic scores it with -- and the evidence for both is spread over four
places that no single file joins:

* ``workflows/rubric_evolution/round_N/`` -- the rubric version produced at round N
  (its ``.md`` text and the edit JSON that says what changed and why),
* ``workflows/round_N/critic_review.json`` -- the critic's per-criterion scores and
  fixes for that round, made with the *previous* round's rubric version,
* ``cpe_evaluations/round_N/<phase>/workflows/round_N/result.json`` -- the judge's
  four dimension scores and their reasons, per problem,
* ``<run>/output/results/interaction_receipt.json`` inside each run -- the actual
  questions asked and the expert's replies.

This script joins them into sheets that can be read top to bottom or pivoted:
``总览`` / ``Rubric演化原因`` / ``Rubric全文`` / ``交互对话`` / ``Judge反馈`` /
``Judge明细`` / ``critic评审``.  It is read-only over the experiment and re-runnable
while the run continues: rounds that have not produced a phase yet are skipped.

    python scripts/export_coevolve_review_workbook.py \
        --exp openclaw_experiments/evolve_exp/claude_fp8_coevolve_from_r10
"""

from __future__ import annotations

import argparse
import json
import re
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]

# Excel's hard limit per cell; longer text is truncated rather than dropped so a
# sheet can never fail to open.
CELL_MAX = 32_767
TRUNCATION_MARK = " …[截断]"

# The critic's own criterion parser (interaction_strategy_critic.rubric_criteria)
# is the format contract for these files: "[20] Name: rule".
CRITERION_RE = re.compile(r"^\[(\d+)\]\s*([^:]+):\s*(.+)$", re.M)
OPERATOR_RE = re.compile(r"Operator\s*(\d+)\s*[:：]\s*([^\n)]+)")

PHASES = (
    "initial_validation_parent_1",
    "initial_train_parent_1",
    "train_parent_1",
    "train_candidate",
    "validation",
)
PHASE_ZH = {
    # The seed's own two phases are how a seeded arm inherits its starting bar;
    # they live under round_0 / round_1 but are not evolved rounds.
    "initial_validation_parent_1": "种子验证(round 0)",
    "initial_train_parent_1": "训练-母代(种子)",
    "train_parent_1": "训练-母代",
    "train_candidate": "训练-候选",
    "validation": "验证",
}

TOP_FIX_RE = re.compile(r"^- \*\*Top fix\*\*:\s*(.+)$", re.M)

DIMENSION_ZH = {
    "analysis_evaluation": "analysis",
    "modeling_rigorousness_evaluation": "rigorousness",
    "practicality_and_scientificity_evaluation": "practicality",
    "result_and_bias_analysis_evaluation": "result_bias",
}


def read_json(path: Path, default=None):
    try:
        with Path(path).open(encoding="utf-8") as handle:
            return json.load(handle)
    except (OSError, ValueError):
        return default


def text(value, limit: int = CELL_MAX) -> str:
    """Cell text: flatten nothing, just cap it so Excel can hold it."""
    if value is None:
        return ""
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False)
    if len(value) > limit:
        return value[: limit - len(TRUNCATION_MARK)] + TRUNCATION_MARK
    return value


def parse_criteria(markdown: str) -> list[dict]:
    return [
        {"max": int(maximum), "name": name.strip(), "rule": rule.strip()}
        for maximum, name, rule in CRITERION_RE.findall(markdown or "")
    ]


def short_id(workflow_id) -> str:
    return str(workflow_id or "").replace("interaction_workflow_", "")


# --------------------------------------------------------------------------- #
# data collection
# --------------------------------------------------------------------------- #
def load_state(exp: Path) -> dict:
    return read_json(exp / "workflows" / "cpe_state.json", {}) or {}


def rubric_versions(exp: Path) -> list[dict]:
    """Every rubric version this run can show: the src default plus its own line."""
    versions = []
    for path in sorted((REPO / "src" / "OpenClaw").glob("interaction_strategy_rubric_v*.md")):
        versions.append(
            {
                "version": f"v{path.stem.rsplit('v', 1)[-1]}",
                "origin": "src 默认（手工维护）",
                "produced_at_round": None,
                "md_path": path,
                "patch": None,
                "prompt_path": None,
            }
        )
    for round_dir in sorted(
        (exp / "workflows" / "rubric_evolution").glob("round_*"),
        key=lambda p: int(p.name.split("_")[1]),
    ):
        round_number = int(round_dir.name.split("_")[1])
        for md_path in sorted(round_dir.glob("interaction_strategy_rubric_v*.md")):
            stem = md_path.stem  # interaction_strategy_rubric_v7
            number = stem.rsplit("v", 1)[-1]
            versions.append(
                {
                    "version": f"v{number}",
                    "origin": f"round {round_number} 末演化产出",
                    "produced_at_round": round_number,
                    "md_path": md_path,
                    "patch": read_json(md_path.with_suffix(".json"), {}),
                    "prompt_path": round_dir / "evolution_prompt.md",
                }
            )
    versions.sort(key=lambda item: int(item["version"][1:]))
    return versions


def rubric_used_by_round(exp: Path) -> dict[int, str]:
    used = {}
    for review_path in sorted((exp / "workflows").glob("round_*/critic_review.json")):
        round_number = int(review_path.parent.name.split("_")[1])
        review = read_json(review_path, {}) or {}
        rubric_id = str(review.get("rubric_id") or "")
        match = re.search(r"_v(\d+)$", rubric_id)
        if match:
            used[round_number] = f"v{match.group(1)}"
    return used


def evidence_scope(prompt_path: Path | None) -> str:
    """Which split the rubric proposal was shown: held-out or training only."""
    if not prompt_path or not prompt_path.is_file():
        return ""
    prompt = prompt_path.read_text(encoding="utf-8", errors="replace")
    if "# 本轮候选在验证集" in prompt:
        return "验证集(held-out)证据"
    if "# 本轮候选在训练集" in prompt:
        return "仅训练集证据"
    return ""


def gate_summary(round_number: int, history: dict[int, dict]) -> str:
    event = history.get(round_number) or {}
    if not event:
        return ""
    parts = [
        f"母代 {event.get('train_pre_utility'):.6f}",
        f"候选 {event.get('candidate_train_net_utility'):.6f}",
        f"训练门 {'过' if event.get('train_accepted') else '未过'}",
    ]
    if event.get("validation_utility") is not None:
        parts.append(f"验证 {float(event['validation_utility']):.6f}")
        parts.append(f"验证门 {'过' if event.get('validation_accepted') else '未过'}")
    else:
        parts.append("验证 未运行")
    return "; ".join(parts)


def phase_results(exp: Path) -> list[dict]:
    """Every stored phase result, oldest round first."""
    collected = []
    for round_dir in sorted(
        (exp / "cpe_evaluations").glob("round_*"),
        key=lambda p: int(p.name.split("_")[1]),
    ):
        round_number = int(round_dir.name.split("_")[1])
        for phase in PHASES:
            path = round_dir / phase / "workflows" / f"round_{round_number}" / "result.json"
            result = read_json(path, None)
            if isinstance(result, dict):
                collected.append(
                    {"round": round_number, "phase": phase, "path": path, "result": result}
                )
    return collected


def repetitions_of(problem_result: dict) -> list[dict]:
    reps = problem_result.get("repetitions")
    if isinstance(reps, list) and reps:
        return reps
    return [problem_result]


def run_receipt(run_dir: str | None) -> dict:
    if not run_dir:
        return {}
    return read_json(Path(run_dir) / "output" / "results" / "interaction_receipt.json", {}) or {}


def top_fix_of(exp: Path, round_number: int) -> str:
    """The review's one-line headline fix; it only exists in the rendered markdown."""
    path = exp / "workflows" / f"round_{round_number}" / "critic_review.md"
    if not path.is_file():
        return ""
    match = TOP_FIX_RE.search(path.read_text(encoding="utf-8", errors="replace"))
    return match.group(1).strip() if match else ""


# --------------------------------------------------------------------------- #
# sheets
# --------------------------------------------------------------------------- #
def sheet_overview(exp: Path) -> list[list]:
    state = load_state(exp)
    history = {int(e["round"]): e for e in state.get("patch_history", [])}
    used = rubric_used_by_round(exp)
    produced = {v["produced_at_round"]: v["version"] for v in rubric_versions(exp) if v["produced_at_round"]}
    patches = {v["produced_at_round"]: v["patch"] for v in rubric_versions(exp) if v["produced_at_round"]}

    rounds = sorted(int(k) for k in (state.get("rounds") or {}))
    # Champion replay: the engine only accepts a candidate that beats the current
    # champion, so the round-by-round champion is the last accepted one.
    seed_bars = [float(r["utility"]) for r in state.get("initial_parent_validation_results", [])]
    champion = ("种子 round 0", max(seed_bars) if seed_bars else float("nan"))
    rows = []
    for round_number in rounds:
        round_state = state["rounds"][str(round_number)]
        event = history.get(round_number, {})
        result = round_state.get("result") or {}
        if event.get("validation_accepted"):
            champion = (f"round {round_number} {short_id(event.get('candidate_workflow_id'))}",
                        float(event.get("validation_utility") or 0.0))
        patch = patches.get(round_number) or {}
        review = read_json(exp / "workflows" / f"round_{round_number}" / "critic_review.json", {}) or {}
        parents = review.get("parents", [])
        top_fix = top_fix_of(exp, round_number)
        critic_mean = parents[0].get("mean_total") if parents else None
        rows.append(
            [
                round_number,
                round_state.get("status") or "进行中",
                used.get(round_number, ""),
                produced.get(round_number, ""),
                (f"{patch.get('op')} / {patch.get('target')}" if patch else ""),
                short_id(event.get("candidate_workflow_id")) or short_id(result.get("workflow_id")),
                result.get("workflow", {}).get("name", ""),
                "; ".join(short_id(x) for x in event.get("training_parent_workflow_ids", [])),
                event.get("train_pre_utility"),
                event.get("candidate_train_net_utility"),
                "过" if event.get("train_accepted") else "未过" if event else "",
                event.get("validation_utility"),
                "过" if event.get("validation_accepted") else ("未过" if event.get("validation_accepted") is False else ""),
                champion[0],
                champion[1],
                critic_mean,
                top_fix,
            ]
        )
    header = [
        "轮次", "状态", "本轮评审用rubric", "本轮末产出的rubric", "产出改动", "候选策略", "候选名称",
        "母代策略", "训练净效用(母代)", "训练净效用(候选)", "训练门", "验证效用(候选)", "验证门",
        "该轮之后的冠军", "冠军验证效用", "母代rubric总分(满分100)", "本轮critic的Top fix",
    ]
    return [header] + rows


def sheet_rubric_reasons(exp: Path) -> list[list]:
    state = load_state(exp)
    history = {int(e["round"]): e for e in state.get("patch_history", [])}
    used = rubric_used_by_round(exp)
    versions = rubric_versions(exp)
    by_version = {v["version"]: v for v in versions}

    rows = []
    for entry in versions:
        if not entry["produced_at_round"]:
            continue
        patch = entry["patch"] or {}
        produced_at = entry["produced_at_round"]
        previous = by_version.get(f"v{int(entry['version'][1:]) - 1}", {})
        previous_criteria = parse_criteria(
            previous["md_path"].read_text(encoding="utf-8", errors="replace")
            if previous.get("md_path")
            else ""
        )
        target = str(patch.get("target") or "")
        before = next((c for c in previous_criteria if c["name"] == target), None)
        new_criteria = patch.get("new_criteria") or []
        after_rule = ""
        for criterion in new_criteria:
            if criterion.get("name") == target or not target:
                after_rule = criterion.get("rule", "")
        applied = patch.get("applied_criteria") or []
        weights = patch.get("weights") or {}
        # Produced at the end of round N, first used by round N+1's critic.
        active_rounds = sorted(r for r, v in used.items() if v == entry["version"])
        rows.append(
            [
                entry["version"],
                produced_at,
                ("".join(str(r) for r in active_rounds)
                 if active_rounds else f"待用（round {produced_at + 1} 起）"),
                patch.get("op", ""),
                target,
                json.dumps(weights, ensure_ascii=False),
                len(applied),
                sum(float(c.get("max", 0)) for c in applied),
                patch.get("evolution_rationale", ""),
                patch.get("predicted_effect", ""),
                gate_summary(produced_at, history),
                evidence_scope(entry.get("prompt_path")),
                before["rule"] if before else "",
                after_rule,
                (len(after_rule) - len(before["rule"])) if before and after_rule else "",
            ]
        )
    header = [
        "rubric版本", "产出轮次", "生效轮次", "改动类型(op)", "改动目标准则", "权重设置", "准则数",
        "权重合计", "为什么这样改（evolution_rationale）", "预期影响（predicted_effect）",
        "产出该版时该轮的门限结果", "证据范围", "改动前该准则全文", "改动后该准则全文", "字符增减",
    ]
    return [header] + rows


def sheet_rubric_text(exp: Path) -> list[list]:
    rows = []
    for entry in rubric_versions(exp):
        body = entry["md_path"].read_text(encoding="utf-8", errors="replace")
        rows.append([entry["version"], entry["origin"], len(body), text(body)])
    return [["rubric版本", "来源", "字符数", "全文"] + [] * 0] + rows


def sheet_dialogues(exp: Path) -> list[list]:
    rows = []
    for item in phase_results(exp):
        result = item["result"]
        workflow_id = result.get("workflow_id")
        for problem in result.get("problem_results", []):
            for repetition in repetitions_of(problem):
                run_dir = repetition.get("run_dir")
                receipt = run_receipt(run_dir)
                questions = receipt.get("questions_asked") or []
                answers = receipt.get("expert_answers") or []
                for index, question in enumerate(questions):
                    reply = answers[index] if index < len(answers) else ""
                    match = OPERATOR_RE.search(str(question))
                    operator = f"{match.group(1)} {match.group(2).strip()}" if match else ""
                    rows.append(
                        [
                            item["round"],
                            PHASE_ZH.get(item["phase"], item["phase"]),
                            repetition.get("problem_id") or problem.get("problem_id"),
                            repetition.get("repetition"),
                            short_id(workflow_id),
                            index + 1,
                            operator,
                            text(question),
                            text(reply),
                            str(Path(run_dir) / "output" / "logs" / "operator_feedback" / "human_expert_dialogue.md")
                            if run_dir
                            else "",
                        ]
                    )
    header = [
        "轮次", "相位", "题目", "重复", "策略", "第几次交换", "算子", "提问", "专家回复", "对话原始文件",
    ]
    return [header] + rows


def sheet_judge_wide(exp: Path) -> list[list]:
    rows = []
    for item in phase_results(exp):
        result = item["result"]
        workflow_id = short_id(result.get("workflow_id"))
        for problem in result.get("problem_results", []):
            scores = problem.get("dimension_scores") or {}
            rows.append(
                [
                    item["round"],
                    PHASE_ZH.get(item["phase"], item["phase"]),
                    problem.get("problem_id"),
                    problem.get("repetition"),
                    problem.get("repetition_count"),
                    problem.get("average_score"),
                    scores.get("analysis_evaluation"),
                    scores.get("modeling_rigorousness_evaluation"),
                    scores.get("practicality_and_scientificity_evaluation"),
                    scores.get("result_and_bias_analysis_evaluation"),
                    (problem.get("interaction_receipt") or {}).get("exchange_count"),
                    problem.get("interaction_cost"),
                    problem.get("interaction_penalty"),
                    problem.get("net_utility"),
                    workflow_id,
                    problem.get("final_report", ""),
                ]
            )
    header = [
        "轮次", "相位", "题目", "重复", "重复数", "报告均分", "analysis", "rigorousness",
        "practicality", "result_bias", "交互次数", "交互成本", "交互惩罚", "净效用", "策略", "报告文件",
    ]
    return [header] + rows


def sheet_judge_long(exp: Path) -> list[list]:
    rows = []
    for item in phase_results(exp):
        for problem in item["result"].get("problem_results", []):
            for repetition in repetitions_of(problem):
                judge = repetition.get("judge_result") or problem.get("judge_result") or {}
                for dimension, analyses in judge.items():
                    if not isinstance(analyses, dict):
                        continue
                    for analysis_name, entry in analyses.items():
                        if not isinstance(entry, dict):
                            continue
                        rows.append(
                            [
                                item["round"],
                                PHASE_ZH.get(item["phase"], item["phase"]),
                                repetition.get("problem_id") or problem.get("problem_id"),
                                repetition.get("repetition"),
                                DIMENSION_ZH.get(dimension, dimension),
                                analysis_name,
                                entry.get("score"),
                                text(entry.get("reason", "")),
                            ]
                        )
    header = ["轮次", "相位", "题目", "重复", "维度", "分析项", "得分", "理由"]
    return [header] + rows


def sheet_critic(exp: Path) -> list[list]:
    rows = []
    for review_path in sorted(
        (exp / "workflows").glob("round_*/critic_review.json"),
        key=lambda p: int(p.parent.name.split("_")[1]),
    ):
        round_number = int(review_path.parent.name.split("_")[1])
        review = read_json(review_path, {}) or {}
        rubric_version = str(review.get("rubric_id") or "")
        top_fix = top_fix_of(exp, round_number)
        for parent in review.get("parents", []):
            for problem in parent.get("per_problem", []):
                for criterion in (problem.get("result") or {}).get("criteria", []):
                    maximum = criterion.get("max") or 0
                    score = criterion.get("score")
                    rows.append(
                        [
                            round_number,
                            rubric_version,
                            parent.get("parent_rank"),
                            short_id(parent.get("workflow_id")),
                            parent.get("mean_total"),
                            parent.get("mean_normalized"),
                            parent.get("net_utility_on_current_training_batch"),
                            problem.get("problem_id"),
                            criterion.get("name"),
                            score,
                            maximum,
                            (score / maximum) if maximum else None,
                            text(criterion.get("justification", "")),
                            text(criterion.get("fix", "")),
                            top_fix,
                        ]
                    )
    header = [
        "轮次", "评审用rubric", "母代序号", "策略", "母代rubric均分", "母代归一化",
        "母代训练净效用", "题目", "准则", "得分", "满分", "归一化",
        "裁决理由", "修复建议", "本轮Top fix",
    ]
    return [header] + rows


def sheet_readme(exp: Path, generated_at: str) -> list[list]:
    return [
        ["本工作簿", "由 scripts/export_coevolve_review_workbook.py 只读导出，可随时重跑刷新"],
        ["实验目录", str(exp)],
        ["生成时间", generated_at],
        ["", ""],
        ["总览", "每轮一行：本轮用哪版 rubric、轮末产出哪版、门限结果、该轮之后的冠军、critic 的 Top fix。种子阶段两行（round 0 验证 / round 1 训练）也在内——被复用的协同臂从种子的 round-0 验证分起算冠军门槛"],
        ["Rubric演化原因", "每版 rubric 一行：改了什么准则、权重、为什么改、预期影响、当时两策略的门限对比、证据范围（验证集 or 仅训练集）、改动前后准则全文"],
        ["Rubric全文", "每版 rubric 全文（含 src 里手工维护的默认版本）"],
        ["交互对话", "每一次专家交换一行：题目/重复/算子/提问/专家回复。对话原文另见各 run 的 output/logs/operator_feedback/human_expert_dialogue.md"],
        ["Judge反馈", "每个 run 一行：MM-Bench 判分四维 + 交互成本 + 净效用（净效用 = 报告均分 - 0.01×交互成本）"],
        ["Judge明细", "judge 的四维×2 项逐条理由（长表，便于按维度透视）"],
        ["critic评审", "当轮 rubric 对母代真实对话的逐准则打分与修复建议（长表）；这是 rubric 影响下一轮策略提议的通道"],
    ]


SHEETS = (
    ("说明", None, [22, 110]),
    ("总览", sheet_overview, [6, 8, 14, 14, 22, 14, 26, 14, 14, 14, 8, 12, 8, 22, 12, 12, 58]),
    ("Rubric演化原因", sheet_rubric_reasons, [10, 10, 12, 10, 20, 22, 7, 8, 70, 60, 46, 16, 70, 70, 9]),
    ("Rubric全文", sheet_rubric_text, [10, 22, 8, 120]),
    ("交互对话", sheet_dialogues, [6, 10, 8, 6, 14, 8, 22, 70, 60, 60]),
    ("Judge反馈", sheet_judge_wide, [6, 10, 8, 6, 7, 10, 10, 12, 12, 12, 9, 10, 10, 10, 14, 55]),
    ("Judge明细", sheet_judge_long, [6, 10, 8, 6, 14, 12, 8, 80]),
    ("critic评审", sheet_critic, [6, 24, 9, 14, 12, 11, 13, 8, 40, 7, 7, 9, 70, 60, 50]),
)


def build(exp: Path, out_path: Path) -> None:
    generated_at = time.strftime("%Y-%m-%d %H:%M:%S")
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for title, builder, widths in SHEETS:
        sheet = workbook.create_sheet(title)
        rows = sheet_readme(exp, generated_at) if builder is None else builder(exp)
        for row in rows:
            sheet.append([text(value) if isinstance(value, str) else value for value in row])
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        if title != "说明":
            sheet.freeze_panes = "A2"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and len(cell.value) > 60:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        print(f"  {title}: {max(sheet.max_row - 1, 0)} 行")
    workbook.save(out_path)
    print(f"写出 {out_path} ({out_path.stat().st_size / 1024:.0f} KB)")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--exp", required=True, help="co-evolution experiment directory")
    parser.add_argument("--out", default=None, help="output .xlsx (default: <exp>/rubric_interaction_judge_review.xlsx)")
    args = parser.parse_args()
    exp = Path(args.exp).resolve()
    if not (exp / "workflows").is_dir():
        raise SystemExit(f"no workflows/ under {exp}")
    out_path = Path(args.out).resolve() if args.out else exp / "rubric_interaction_judge_review.xlsx"
    build(exp, out_path)


if __name__ == "__main__":
    main()
