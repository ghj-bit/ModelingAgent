"""Compare the two four-dimension Judge passes of a workflow experiment."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path
from typing import Any


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_EXPERIMENT = (
    REPO_ROOT
    / "openclaw_experiments"
    / "interaction_workflow_substantive_20260909_214301"
)
DIMENSIONS = (
    "structural_coherency",
    "scoring_decomposition",
    "modeling_groundedness",
    "analysis_groundedness",
)
LABELS = {
    "structural_coherency": "结构完整度",
    "scoring_decomposition": "任务覆盖度",
    "modeling_groundedness": "建模扎实度",
    "analysis_groundedness": "分析扎实度",
}


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def load_rows(experiment: Path) -> list[dict[str, Any]]:
    first_path = experiment / "workflows" / "results.json"
    second_path = (
        experiment
        / "workflows"
        / "deepseek_v4_pro_six_dimension_evaluation"
        / "results.json"
    )
    first = read_json(first_path)
    second = read_json(second_path)
    first_by_key = {
        (int(node["round"]), result["problem_id"]): result
        for node in first
        for result in node.get("problem_results", [])
    }
    second_by_key = {
        (int(item["round"]), item["problem_id"]): item
        for item in second.get("reports", [])
        if int(item.get("repetition", 1)) == 1
    }
    if set(first_by_key) != set(second_by_key):
        missing_second = sorted(set(first_by_key) - set(second_by_key))
        missing_first = sorted(set(second_by_key) - set(first_by_key))
        raise ValueError(
            f"Judge result keys do not match; missing second={missing_second}, "
            f"missing first={missing_first}"
        )

    rows = []
    for round_number, problem_id in sorted(first_by_key):
        first_item = first_by_key[(round_number, problem_id)]
        second_item = second_by_key[(round_number, problem_id)]
        first_scores = first_item.get("dimension_scores", {})
        second_scores = second_item.get("dimension_scores", {})
        if any(
            not isinstance(scores.get(name), (int, float))
            for scores in (first_scores, second_scores)
            for name in DIMENSIONS
        ):
            raise ValueError(f"Incomplete four-dimension scores: {(round_number, problem_id)}")
        comparisons = {
            name: {
                "first": float(first_scores[name]),
                "second": float(second_scores[name]),
                "delta": float(second_scores[name]) - float(first_scores[name]),
                "absolute_delta": abs(
                    float(second_scores[name]) - float(first_scores[name])
                ),
            }
            for name in DIMENSIONS
        }
        first_average = sum(x["first"] for x in comparisons.values()) / 4
        second_average = sum(x["second"] for x in comparisons.values()) / 4
        absolute_deltas = [x["absolute_delta"] for x in comparisons.values()]
        rows.append(
            {
                "round": round_number,
                "problem_id": problem_id,
                "comparisons": comparisons,
                "first_average": first_average,
                "second_average": second_average,
                "average_delta": second_average - first_average,
                "mean_absolute_delta": sum(absolute_deltas) / 4,
                "max_absolute_delta": max(absolute_deltas),
                "first_judge_result": first_item.get("judge_result", ""),
                "second_judge_result": second_item.get("judge_result", ""),
            }
        )
    return rows


def export_excel(rows: list[dict[str, Any]], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    detail = workbook.active
    detail.title = "逐题对比"
    headers = [
        "轮次",
        "题目",
        "首次四维平均",
        "二次四维平均",
        "平均分差(二次-首次)",
        "四维平均绝对差",
        "最大维度绝对差",
    ]
    for name in DIMENSIONS:
        label = LABELS[name]
        headers.extend([f"{label}-首次", f"{label}-二次", f"{label}-差值", f"{label}-绝对差"])
    headers.extend(["首次Judge文件", "二次Judge文件"])
    detail.append(headers)
    for row in rows:
        values = [
            row["round"],
            row["problem_id"],
            row["first_average"],
            row["second_average"],
            row["average_delta"],
            row["mean_absolute_delta"],
            row["max_absolute_delta"],
        ]
        for name in DIMENSIONS:
            comparison = row["comparisons"][name]
            values.extend(
                [
                    comparison["first"],
                    comparison["second"],
                    comparison["delta"],
                    comparison["absolute_delta"],
                ]
            )
        values.extend([row["first_judge_result"], row["second_judge_result"]])
        detail.append(values)

    dimension_detail = workbook.create_sheet("逐维明细")
    dimension_detail.append(
        ["轮次", "题目", "维度", "首次分数", "二次分数", "差值(二次-首次)", "绝对差"]
    )
    for row in rows:
        for name in DIMENSIONS:
            comparison = row["comparisons"][name]
            dimension_detail.append(
                [
                    row["round"],
                    row["problem_id"],
                    LABELS[name],
                    comparison["first"],
                    comparison["second"],
                    comparison["delta"],
                    comparison["absolute_delta"],
                ]
            )

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["problem_id"]].append(row)
    summary = workbook.create_sheet("题目汇总")
    summary_headers = [
        "题目",
        "轮次数",
        "首次四维均分",
        "二次四维均分",
        "平均分差",
        "四维平均绝对差",
        "单项最大绝对差",
        *[f"{LABELS[name]}平均绝对差" for name in DIMENSIONS],
    ]
    summary.append(summary_headers)
    for problem_id, items in sorted(grouped.items()):
        summary.append(
            [
                problem_id,
                len(items),
                sum(x["first_average"] for x in items) / len(items),
                sum(x["second_average"] for x in items) / len(items),
                sum(x["average_delta"] for x in items) / len(items),
                sum(x["mean_absolute_delta"] for x in items) / len(items),
                max(x["max_absolute_delta"] for x in items),
                *[
                    sum(x["comparisons"][name]["absolute_delta"] for x in items)
                    / len(items)
                    for name in DIMENSIONS
                ],
            ]
        )

    notes = workbook.create_sheet("说明")
    notes.append(["项目", "说明"])
    notes.append(["首次Judge", "实验执行阶段保存于 workflows/results.json 的四维评分"])
    notes.append(["二次Judge", "六维复评结果中的相同四个维度"])
    notes.append(["差值", "二次分数减首次分数；正数表示二次评分更高"])
    notes.append(["绝对差参考", "≤0.05较小；0.05–0.10中等；>0.10较大"])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    absolute_scale = ColorScaleRule(
        start_type="num",
        start_value=0,
        start_color="63BE7B",
        mid_type="num",
        mid_value=0.1,
        mid_color="FFEB84",
        end_type="num",
        end_value=0.25,
        end_color="F8696B",
    )
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_index, cells in enumerate(sheet.columns, 1):
            width = min(max(max(len(str(cell.value or "")) for cell in cells) + 2, 10), 60)
            sheet.column_dimensions[get_column_letter(column_index)].width = width

    for sheet in (detail, dimension_detail, summary):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    if detail.max_row > 1:
        detail.conditional_formatting.add(f"F2:G{detail.max_row}", absolute_scale)
        for column in (10, 14, 18, 22):
            letter = get_column_letter(column)
            detail.conditional_formatting.add(
                f"{letter}2:{letter}{detail.max_row}", absolute_scale
            )
    if dimension_detail.max_row > 1:
        dimension_detail.conditional_formatting.add(
            f"G2:G{dimension_detail.max_row}", absolute_scale
        )
    if summary.max_row > 1:
        summary.conditional_formatting.add(f"F2:K{summary.max_row}", absolute_scale)

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    workbook.save(temporary)
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--experiment", type=Path, default=DEFAULT_EXPERIMENT)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    experiment = args.experiment.resolve()
    output = (
        args.output.resolve()
        if args.output
        else experiment / "workflows" / "two_judge_four_dimension_comparison.xlsx"
    )
    rows = load_rows(experiment)
    export_excel(rows, output)
    overall_mean_absolute = sum(x["mean_absolute_delta"] for x in rows) / len(rows)
    overall_maximum = max(x["max_absolute_delta"] for x in rows)
    print(f"Compared reports: {len(rows)}")
    print(f"Mean absolute dimension difference: {overall_mean_absolute:.4f}")
    print(f"Maximum absolute dimension difference: {overall_maximum:.4f}")
    print(f"Excel: {output}")


if __name__ == "__main__":
    main()
