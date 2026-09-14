"""Compare four shared dimensions in two workflow-experiment Excel files."""

from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_EXPERIMENT = (
    REPO_ROOT
    / "openclaw_experiments"
    / "interaction_workflow_substantive_20260909_214301"
)
DIMENSIONS = (
    "structural_coherency",
    "scoring_decomposition",
    "modeling_groundedness",
    "analysis_groundedness",
)
LABELS = {
    "structural_coherency": "结构完整度",
    "scoring_decomposition": "任务覆盖度",
    "modeling_groundedness": "建模扎实度",
    "analysis_groundedness": "分析扎实度",
}


def header_map(sheet) -> dict[str, int]:
    return {str(cell.value): cell.column for cell in sheet[1] if cell.value is not None}


def load_original(path: Path) -> dict[tuple[int, str], dict[str, float]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["长表_便于透视"]
    headers = header_map(sheet)
    required = ("轮次", "Problem ID", "Dimension Key", "平均得分")
    if any(name not in headers for name in required):
        raise ValueError(f"Original score workbook has unexpected headers: {path}")
    scores: dict[tuple[int, str], dict[str, float]] = defaultdict(dict)
    for row in sheet.iter_rows(min_row=2, values_only=True):
        dimension = row[headers["Dimension Key"] - 1]
        if dimension not in DIMENSIONS:
            continue
        key = (int(row[headers["轮次"] - 1]), str(row[headers["Problem ID"] - 1]))
        scores[key][str(dimension)] = float(row[headers["平均得分"] - 1])
    return dict(scores)


def load_repeat(path: Path) -> dict[tuple[int, str], dict[str, float]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["轮次题目汇总"]
    headers = header_map(sheet)
    required = ("轮次", "题目", *[LABELS[name] for name in DIMENSIONS])
    if any(name not in headers for name in required):
        raise ValueError(f"Repeat score workbook has unexpected headers: {path}")
    scores = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        key = (int(row[headers["轮次"] - 1]), str(row[headers["题目"] - 1]))
        scores[key] = {
            name: float(row[headers[LABELS[name]] - 1]) for name in DIMENSIONS
        }
    return scores


def comparison_rows(
    original: dict[tuple[int, str], dict[str, float]],
    repeat: dict[tuple[int, str], dict[str, float]],
) -> list[dict[str, Any]]:
    if set(original) != set(repeat):
        raise ValueError(
            "Workbook keys do not match: "
            f"only original={sorted(set(original) - set(repeat))}; "
            f"only repeat={sorted(set(repeat) - set(original))}"
        )
    rows = []
    for round_number, problem_id in sorted(original):
        if set(original[(round_number, problem_id)]) != set(DIMENSIONS):
            raise ValueError(f"Original workbook has incomplete scores: {(round_number, problem_id)}")
        comparisons = {}
        for name in DIMENSIONS:
            first = original[(round_number, problem_id)][name]
            second = repeat[(round_number, problem_id)][name]
            comparisons[name] = {
                "first": first,
                "second": second,
                "delta": second - first,
                "absolute_delta": abs(second - first),
            }
        first_average = sum(item["first"] for item in comparisons.values()) / 4
        second_average = sum(item["second"] for item in comparisons.values()) / 4
        absolute = [item["absolute_delta"] for item in comparisons.values()]
        rows.append(
            {
                "round": round_number,
                "problem_id": problem_id,
                "comparisons": comparisons,
                "first_average": first_average,
                "second_average": second_average,
                "average_delta": second_average - first_average,
                "mean_absolute_delta": sum(absolute) / 4,
                "max_absolute_delta": max(absolute),
            }
        )
    return rows


def append_aggregate_sheet(workbook: Workbook, title: str, groups: dict[Any, list[dict]]) -> None:
    sheet = workbook.create_sheet(title)
    group_label = "题目" if title == "题目汇总" else "轮次"
    sheet.append(
        [
            group_label,
            "样本数",
            "首次四维均分",
            "复评四维均分",
            "平均分差",
            "四维平均绝对差",
            "单项最大绝对差",
            *[f"{LABELS[name]}平均绝对差" for name in DIMENSIONS],
        ]
    )
    for group, items in sorted(groups.items()):
        sheet.append(
            [
                group,
                len(items),
                sum(item["first_average"] for item in items) / len(items),
                sum(item["second_average"] for item in items) / len(items),
                sum(item["average_delta"] for item in items) / len(items),
                sum(item["mean_absolute_delta"] for item in items) / len(items),
                max(item["max_absolute_delta"] for item in items),
                *[
                    sum(item["comparisons"][name]["absolute_delta"] for item in items)
                    / len(items)
                    for name in DIMENSIONS
                ],
            ]
        )


def export_excel(rows: list[dict[str, Any]], output: Path, original: Path, repeat: Path) -> None:
    workbook = Workbook()
    wide = workbook.active
    wide.title = "逐题对比"
    headers = [
        "轮次",
        "题目",
        "首次四维平均",
        "复评四维平均",
        "平均分差(复评-首次)",
        "四维平均绝对差",
        "最大维度绝对差",
    ]
    for name in DIMENSIONS:
        label = LABELS[name]
        headers.extend([f"{label}-首次", f"{label}-复评", f"{label}-差值", f"{label}-绝对差"])
    wide.append(headers)
    for item in rows:
        values = [
            item["round"],
            item["problem_id"],
            item["first_average"],
            item["second_average"],
            item["average_delta"],
            item["mean_absolute_delta"],
            item["max_absolute_delta"],
        ]
        for name in DIMENSIONS:
            value = item["comparisons"][name]
            values.extend([value["first"], value["second"], value["delta"], value["absolute_delta"]])
        wide.append(values)

    long_sheet = workbook.create_sheet("逐维明细")
    long_sheet.append(["轮次", "题目", "维度", "首次分数", "复评分数", "差值", "绝对差"])
    for item in rows:
        for name in DIMENSIONS:
            value = item["comparisons"][name]
            long_sheet.append(
                [item["round"], item["problem_id"], LABELS[name], value["first"], value["second"], value["delta"], value["absolute_delta"]]
            )

    by_problem: dict[str, list[dict]] = defaultdict(list)
    by_round: dict[int, list[dict]] = defaultdict(list)
    for item in rows:
        by_problem[item["problem_id"]].append(item)
        by_round[item["round"]].append(item)
    append_aggregate_sheet(workbook, "题目汇总", by_problem)
    append_aggregate_sheet(workbook, "轮次汇总", by_round)

    notes = workbook.create_sheet("说明")
    notes.append(["项目", "说明"])
    notes.append(["首次评分文件", str(original)])
    notes.append(["复评文件", str(repeat)])
    notes.append(["比较范围", f"{len(rows)}个轮次题目配对 × 4维 = {len(rows) * 4}个配对分数"])
    notes.append(["差值", "复评分数减首次分数；正值表示复评更高"])
    notes.append(["稳定性参考", "绝对差≤0.05较稳定；0.05–0.10中等；>0.10差异较大。仅为人工检查参考；报告变化的对比不能单独视作随机评分稳定性。"])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    scale = ColorScaleRule(
        start_type="num", start_value=0, start_color="63BE7B",
        mid_type="num", mid_value=0.1, mid_color="FFEB84",
        end_type="num", end_value=0.25, end_color="F8696B",
    )
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for index, cells in enumerate(sheet.columns, 1):
            width = min(max(max(len(str(cell.value or "")) for cell in cells) + 2, 10), 65)
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    wide.conditional_formatting.add(f"F2:G{wide.max_row}", scale)
    for column in (11, 15, 19, 23):
        letter = get_column_letter(column)
        wide.conditional_formatting.add(f"{letter}2:{letter}{wide.max_row}", scale)
    long_sheet.conditional_formatting.add(f"G2:G{long_sheet.max_row}", scale)
    for title in ("题目汇总", "轮次汇总"):
        sheet = workbook[title]
        sheet.conditional_formatting.add(f"F2:K{sheet.max_row}", scale)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    workbook.save(temporary)
    temporary.replace(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--experiment", type=Path, default=DEFAULT_EXPERIMENT)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    experiment = args.experiment.resolve()
    original = experiment / "workflows" / "round_problem_dimension_scores.xlsx"
    repeat = experiment / "workflows" / "deepseek_v4_flash_six_dimension_evaluation" / "six_dimension_scores.xlsx"
    output = args.output.resolve() if args.output else experiment / "workflows" / "four_dimension_judge_stability_comparison.xlsx"
    rows = comparison_rows(load_original(original), load_repeat(repeat))
    export_excel(rows, output, original, repeat)
    mean_absolute = sum(item["mean_absolute_delta"] for item in rows) / len(rows)
    maximum = max(item["max_absolute_delta"] for item in rows)
    print(f"Compared pairs: {len(rows) * len(DIMENSIONS)}")
    print(f"Mean absolute difference: {mean_absolute:.4f}")
    print(f"Maximum absolute difference: {maximum:.4f}")
    print(f"Excel: {output}")


if __name__ == "__main__":
    main()
