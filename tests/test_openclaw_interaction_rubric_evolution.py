from pathlib import Path
import os
import sys
import tempfile
import time
import unittest
from types import SimpleNamespace
from unittest.mock import patch


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from src.OpenClaw import baseline
from src.OpenClaw import run_interaction_rubric_evolution as interaction


class InteractionRubricEvolutionTests(unittest.TestCase):
    def test_default_expert_model_is_deepseek_v4_pro(self):
        with patch.object(sys, "argv", ["run_interaction_rubric_evolution.py"]):
            args = interaction.parse_args()

        self.assertEqual(args.expert_model, "deepseek-v4-pro")
        self.assertTrue(args.skip_interaction_receipt_validation)

    def test_initial_prompt_injects_interaction_before_workflow_and_report(self):
        base_prompt = baseline.BASELINE_PROMPT_TEMPLATE_PATH.read_text(encoding="utf-8")
        prompt = interaction.build_prompt(base_prompt, interaction.INITIAL_RUBRIC)
        normalized_prompt = " ".join(prompt.split())

        self.assertLess(
            prompt.index("## Human Expert Interaction"),
            prompt.index("## Required Workflow"),
        )
        self.assertLess(
            prompt.index("## Human Expert Interaction"),
            prompt.index("## Final Report Contract"),
        )
        self.assertIn("Good example:", prompt)
        self.assertIn("Poor example:", prompt)
        self.assertIn("during Required Workflow steps 1-5", prompt)
        self.assertIn("human_expert_dialogue.md", prompt)
        self.assertNotIn("react_verifier_interaction.md", prompt)
        self.assertNotIn("interaction_evaluation.md", prompt)
        self.assertIn("blocking tool", prompt)
        self.assertIn("continue in this same run", normalized_prompt)
        self.assertNotIn("stop the current turn", prompt)
        self.assertIn("Never delegate calculations", prompt)
        self.assertIn("modeling agent owns all technical work", normalized_prompt)
        self.assertNotIn("post-implementation check", prompt)
        self.assertIn("interaction_stage` from 1 to 5", prompt)
        self.assertIn(
            f"`{interaction.INITIAL_RUBRIC['rubric_id']}`; do not use the rubric name",
            normalized_prompt,
        )
        self.assertIn("## Workflow Summary Contract", prompt)
        self.assertIn("step_01_problem_understanding.md", prompt)
        self.assertIn("step_02_modeling_assumptions.md", prompt)
        self.assertIn("external_data.md", prompt)
        self.assertIn(
            "never directly adopt an expert-supplied number or technical parameter",
            normalized_prompt.lower(),
        )
        self.assertIn("compare exactly two agent-derived technical", prompt)
        self.assertIn("third resolving scenario only", prompt)
        self.assertIn("roughly 2-3 minutes", prompt)
        self.assertIn("pure value preference", prompt)
        self.assertIn("at most 1-2 authoritative sources", prompt)
        self.assertIn("interaction_decision.json", prompt)
        self.assertNotIn('```json', prompt)
        self.assertNotIn("Before starting a later step, read", prompt)
        self.assertIn("ADOPTED_VALUE", prompt)
        self.assertIn("ADAPTED_TECHNICALLY", prompt)
        self.assertIn("REJECTED_AFTER_VALIDATION", prompt)
        self.assertIn("step_04_implementation.md", prompt)
        self.assertIn("step_05_validation_analysis.md", prompt)
        self.assertIn("Start step 6 only after these five", prompt)
        self.assertLess(
            prompt.index("## Workflow Summary Contract"),
            prompt.index("## Final Report Contract"),
        )

    def test_initial_rubric_and_validation_problem_contract(self):
        interaction.validate_rubric(interaction.INITIAL_RUBRIC)
        self.assertEqual(len(interaction.INITIAL_RUBRIC["criteria"]), 1)
        self.assertEqual(
            interaction.INITIAL_RUBRIC["rubric_id"],
            "interaction_initial_minimal_v3",
        )

        self.assertNotIn("purpose", interaction.INITIAL_RUBRIC)
        self.assertNotIn("success_test", interaction.INITIAL_RUBRIC)
        self.assertEqual(
            interaction.VALIDATION_PROBLEMS,
            (
                "2013_Bank_Service_Problem",
                "2025_Managing_Sustainable_Tourism",
                "2003_Aviation_Baggage_Screening",
            ),
        )

    def test_execution_prompt_injects_root_without_exposing_it_to_optimizer(self):
        evolved = {
            "rubric_id": "interaction_evolved",
            "name": "Extended rubric",
            "attention_budget": "Use one exchange for a qualitative decision.",
            "criteria": [{
            "name": "Preference framing",
            "rule": "Ask which stakeholder preference should govern one unresolved qualitative tradeoff.",
            "positive_example": "Present two qualitative policy priorities and request one choice.",
            "negative_example": "Ask the expert to calculate which policy has the larger score.",
            }],
        }
        rendered = interaction.format_execution_rubrics(evolved)
        self.assertIn("Fixed initial rubric", rendered)
        self.assertIn(interaction.INITIAL_RUBRIC["criteria"][0]["rule"], rendered)
        self.assertIn(evolved["criteria"][0]["rule"], rendered)
        self.assertIsNone(
            interaction.optimizer_visible_rubric(interaction.INITIAL_RUBRIC)
        )
        self.assertEqual(interaction.optimizer_visible_rubric(evolved), evolved)
        self.assertNotIn("innovativeness", interaction.EVALUATION_DIMENSIONS)
        self.assertNotIn("data_groundedness", interaction.EVALUATION_DIMENSIONS)
        self.assertEqual(len(interaction.EVALUATION_DIMENSIONS), 4)

    def test_judge_report_uses_short_workspace_parent(self):
        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary) / ("nested_" * 20)
            result = {
                "run_dir": run_dir,
                "final_report": run_dir / "output/results/solution_report.md",
            }
            args = SimpleNamespace(judge_repeats=1, judge_concurrency=1)
            payload = {
                "average_scores": {"round_report": {
                    "trial_count": 1,
                    "average_dimension_scores": {
                        name: 0.8 for name in interaction.EVALUATION_DIMENSIONS
                    },
                }},
                "evaluations": {"report": [{"raw_result": "judge.json"}]},
            }
            with patch.object(
                interaction.run_judge_stability,
                "evaluate_reports_repeated",
                return_value=payload,
            ) as evaluate:
                interaction.judge_report("problem", result, 2, run_dir, args)

            workspace_parent = evaluate.call_args.args[5]
            self.assertEqual(
                workspace_parent,
                interaction.short_judge_workspace_parent(run_dir),
            )
            self.assertTrue(str(workspace_parent).startswith(str(REPO_ROOT)))
            self.assertLess(
                len(
                    str(
                        workspace_parent / ".judge" / "j_test"
                        / "final_submission" / "solution_report.md"
                    )
                ),
                260,
            )

    def test_dimension_aggregation_uses_both_validation_problems(self):
        scores = interaction.average_dimensions(
            [
                {"dimension_scores": {"modeling": 0.6, "analysis": 0.8}},
                {"dimension_scores": {"modeling": 0.8, "analysis": 0.6}},
            ]
        )
        self.assertEqual(scores, {"analysis": 0.7, "modeling": 0.7})

    def test_score_summary_contains_round_problem_and_four_dimensions(self):
        dimensions = {
            "analysis_groundedness": 0.6,
            "modeling_groundedness": 0.7,
            "scoring_decomposition": 0.8,
            "structural_coherency": 0.9,
        }
        summary = interaction.build_score_summary(
            Path("experiment"),
            [
                {
                    "round": 1,
                    "rubric_id": "rubric_1",
                    "problem_results": [
                        {
                            "problem_id": "problem_1",
                            "dimension_scores": dimensions,
                            "average_score": 0.0,
                        }
                    ],
                }
            ],
        )

        self.assertEqual(summary["best_round"]["round"], 1)
        problem = summary["rounds"][0]["problems"][0]
        self.assertEqual(problem["dimension_scores"], dimensions)
        self.assertAlmostEqual(problem["four_dimension_average"], 0.75)
        self.assertNotIn("innovativeness", summary["dimensions"])
        self.assertNotIn("data_groundedness", summary["dimensions"])

    def test_score_workbook_groups_dimensions_by_problem(self):
        from openpyxl import load_workbook

        dimensions = {
            "analysis_groundedness": 0.6,
            "modeling_groundedness": 0.7,
            "scoring_decomposition": 0.8,
            "structural_coherency": 0.9,
        }
        results = [
            {
                "round": 1,
                "rubric_id": "rubric_1",
                "rubric": {"name": "Test rubric"},
                "parent_round": None,
                "problem_results": [
                    {
                        "problem_id": problem_id,
                        "dimension_scores": dimensions,
                        "average_score": 0.0,
                    }
                    for problem_id in interaction.VALIDATION_PROBLEMS
                ],
            }
        ]
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "scores.xlsx"
            interaction.write_score_workbook(
                output, Path(temporary) / "experiment", results
            )
            workbook = load_workbook(output, data_only=True)

        self.assertEqual(
            workbook.sheetnames,
            [
                "总览_按题目",
                "题目_银行服务",
                "题目_可持续旅游",
                "题目_行李安检",
                "长表_便于透视",
                "说明",
            ],
        )
        overview = workbook["总览_按题目"]
        self.assertEqual(overview["A3"].value, 1)
        self.assertAlmostEqual(overview["D3"].value, 0.75)
        self.assertIn("2013_Bank_Service_Problem", overview["F1"].value)
        self.assertEqual(overview["F2"].value, "四维平均")
        self.assertAlmostEqual(overview["F3"].value, 0.75)
        bank = workbook["题目_银行服务"]
        self.assertEqual(bank["A2"].value, 1)
        self.assertAlmostEqual(bank["D2"].value, 0.75)
        self.assertEqual(workbook["长表_便于透视"].max_row, 16)

    def test_interaction_workbook_contains_rubric_and_concrete_dialogue(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            problem_results = []
            for problem_id in interaction.VALIDATION_PROBLEMS:
                run_dir = root / problem_id
                dialogue = (
                    run_dir
                    / "output"
                    / "logs"
                    / "operator_feedback"
                    / "human_expert_dialogue.md"
                )
                dialogue.parent.mkdir(parents=True)
                dialogue.write_text(
                    f"Question for {problem_id}\nExpert qualitative answer",
                    encoding="utf-8",
                )
                problem_results.append(
                    {
                        "problem_id": problem_id,
                        "run_dir": str(run_dir),
                        "repetition": 1,
                        "interaction_receipt": {
                            "interaction_stage": 2,
                            "exchange_count": 1,
                            "questions_asked": [f"Question for {problem_id}"],
                            "adoption_decisions": [
                                {
                                    "expert_recommendation": "Prioritize safety.",
                                    "input_type": "value_preference",
                                    "status": "ADOPTED_VALUE",
                                    "adopted_scope": "Safety priority only.",
                                    "technical_translations": [
                                        {"name": "Alternative A", "result": "Feasible"},
                                        {"name": "Alternative B", "result": "Safer"},
                                    ],
                                    "validation": {
                                        "method": "Compared two alternatives.",
                                        "comparison_result": "Alternative B was safer.",
                                        "external_data": {
                                            "search_performed": True,
                                            "sources": [
                                                {
                                                    "source": "Empirical source",
                                                    "empirical_use": "Validated reliability",
                                                    "evidence_role": "validation",
                                                }
                                            ],
                                            "data_gap": "",
                                            "claim_limit": "",
                                        },
                                        "evidence_files": [
                                            "results/validation.json"
                                        ],
                                    },
                                }
                            ],
                            "changed_files": ["results/assumptions.md"],
                            "validation_files": ["results/validation.json"],
                        },
                    }
                )
            results = [
                {
                    "round": 1,
                    "rubric_id": interaction.INITIAL_RUBRIC["rubric_id"],
                    "rubric": interaction.INITIAL_RUBRIC,
                    "parent_round": None,
                    "utility": 0.8,
                    "problem_results": problem_results,
                }
            ]
            output = root / "interactions.xlsx"
            interaction.write_interaction_workbook(output, root, results)
            workbook = load_workbook(output, data_only=True)

        self.assertEqual(
            workbook.sheetnames,
            [
                "交互总览_按题目",
                "Rubric演化",
                "题目_银行服务",
                "题目_可持续旅游",
                "题目_行李安检",
                "说明",
            ],
        )
        overview = workbook["交互总览_按题目"]
        self.assertEqual(overview["A3"].value, 1)
        self.assertIn("2013_Bank_Service_Problem", overview["D1"].value)
        self.assertIn("Question for", overview["F3"].value)
        self.assertIn("Expert qualitative answer", overview["G3"].value)
        rubric_sheet = workbook["Rubric演化"]
        self.assertIn("One useful decision", rubric_sheet["G2"].value)
        bank = workbook["题目_银行服务"]
        self.assertEqual(bank["I2"].value, 2)
        self.assertEqual(bank["J2"].value, 1)
        self.assertIn("ADOPTED_VALUE", bank["M2"].value)

    def test_interaction_validation_enforces_structured_verified_adoption(self):
        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary)
            output = run_dir / "output"
            dialogue = output / "logs/operator_feedback/human_expert_dialogue.md"
            external_data = output / "data/external_data.md"
            validation_file = output / "results/validation.json"
            receipt_path = output / "results/interaction_receipt.json"
            report = output / "results/solution_report.md"
            for path, content in (
                (dialogue, "expert dialogue"),
                (external_data, "external source and its validation use"),
                (validation_file, '{"validated": true}'),
            ):
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(content, encoding="utf-8")
            receipt = {
                "rubric_id": interaction.INITIAL_RUBRIC["rubric_id"],
                "interaction_stage": 2,
                "questions_asked": ["Which qualitative value should dominate?"],
                "exchange_count": 1,
                "adoption_decisions": [
                    {
                        "core_decision": "Prioritize safety over cost.",
                        "input_type": "value_preference",
                        "status": "ADOPTED_VALUE",
                        "adopted_scope": "Only the qualitative risk priority.",
                        "technical_translations": [
                            {
                                "name": "99 percent service",
                                "evidence_file": "results/validation.json",
                            },
                            {
                                "name": "99.9 percent service",
                                "evidence_file": "results/validation.json",
                            },
                        ],
                        "comparison_after_two": "stable",
                        "external_validation": {
                            "required": False,
                            "sources": [],
                            "data_gap": "",
                            "claim_limit": "",
                        },
                    }
                ],
                "changed_files": ["results/solution_report.md"],
                "validation_files": ["results/validation.json"],
                "interaction_complete": True,
            }
            interaction.workflow_evolution.write_json(receipt_path, receipt)
            report.write_text("final report after validation", encoding="utf-8")
            result = {"run_dir": run_dir, "final_report": report}

            validated = interaction.validate_interaction_run(
                result, interaction.INITIAL_RUBRIC
            )

        self.assertEqual(
            validated["adoption_decisions"][0]["status"], "ADOPTED_VALUE"
        )

    def test_interaction_validation_rejects_unverified_technical_adoption(self):
        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary)
            output = run_dir / "output"
            paths = {
                "dialogue": output
                / "logs/operator_feedback/human_expert_dialogue.md",
                "external": output / "data/external_data.md",
                "validation": output / "results/validation.json",
                "receipt": output / "results/interaction_receipt.json",
                "report": output / "results/solution_report.md",
            }
            for key in ("dialogue", "external", "validation"):
                paths[key].parent.mkdir(parents=True, exist_ok=True)
                paths[key].write_text(key, encoding="utf-8")
            receipt = {
                "rubric_id": interaction.INITIAL_RUBRIC["rubric_id"],
                "interaction_stage": 2,
                "questions_asked": ["Choose a frame."],
                "exchange_count": 1,
                "adoption_decisions": [
                    {
                        "expert_recommendation": "Use the expert's P90 value.",
                        "input_type": "technical_suggestion",
                        "status": "ADOPTED_VALUE",
                        "adopted_scope": "Directly used P90.",
                        "technical_translations": [
                            {"name": "P90", "result": "adopted"}
                        ],
                        "validation": {
                            "method": "No independent comparison.",
                            "comparison_result": "Accepted directly.",
                            "external_data": {
                                "search_performed": False,
                                "sources": [],
                                "data_gap": "",
                                "claim_limit": "",
                            },
                            "evidence_files": ["results/validation.json"],
                        },
                    }
                ],
                "changed_files": [],
                "validation_files": ["results/validation.json"],
                "interaction_complete": True,
            }
            interaction.workflow_evolution.write_json(paths["receipt"], receipt)
            paths["report"].write_text("report", encoding="utf-8")

            with self.assertRaisesRegex(
                RuntimeError,
                "ADOPTED_VALUE only for value_preference",
            ):
                interaction.validate_interaction_run(
                    {"run_dir": run_dir, "final_report": paths["report"]},
                    interaction.INITIAL_RUBRIC,
                )

    def test_run_validation_problem_can_skip_receipt_validation_for_scoring(self):
        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary)
            report = run_dir / "output/results/solution_report.md"
            receipt_path = run_dir / "output/results/interaction_receipt.json"
            feedback = run_dir / "output/logs/operator_feedback"
            report.parent.mkdir(parents=True, exist_ok=True)
            report.write_text("existing final report", encoding="utf-8")
            interaction.workflow_evolution.write_json(
                feedback / "expert_request_1.json",
                {
                    "question": "Which qualitative priority should govern?",
                    "interaction_stage": 2,
                    "rubric_id": "human-readable-name-is-tolerated",
                },
            )
            interaction.workflow_evolution.write_json(
                feedback / "expert_reply_1.json",
                {"ok": True, "answer": "Prefer safety."},
            )
            (feedback / "human_expert_dialogue.md").write_text(
                "question and expert answer", encoding="utf-8"
            )
            interaction.workflow_evolution.write_json(
                receipt_path, {"rubric_id": "intentionally-incomplete"}
            )
            prepared = {
                "problem_id": interaction.VALIDATION_PROBLEMS[0],
                "repetition": 1,
                "run_dir": run_dir,
                "prompt": run_dir / "prompt.md",
                "final_report": report,
                "recovered": True,
            }
            args = SimpleNamespace(skip_interaction_receipt_validation=True)
            judged = {
                "average_score": 0.8,
                "dimension_scores": {"structural_coherency": 0.8},
                "judge_repeats": 1,
                "judge_stability_result": "stability.json",
                "judge_result": "judge.json",
            }

            with patch.object(
                interaction, "validate_interaction_run"
            ) as validate_receipt, patch.object(
                interaction, "judge_report", return_value=judged
            ) as judge:
                result = interaction.run_validation_problem(
                    prepared,
                    interaction.INITIAL_RUBRIC,
                    1,
                    run_dir,
                    args,
                )

            validate_receipt.assert_not_called()
            judge.assert_called_once()
            self.assertEqual(result["average_score"], 0.8)
            self.assertEqual(
                result["interaction_receipt"]["rubric_id"],
                "intentionally-incomplete",
            )

    def test_python_finalizes_compact_decision_without_repeating_results(self):
        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary)
            output = run_dir / "output"
            feedback = output / "logs/operator_feedback"
            evidence = output / "results/comparison.json"
            for path in (feedback, evidence.parent):
                path.mkdir(parents=True, exist_ok=True)
            interaction.workflow_evolution.write_json(
                feedback / "expert_request_1.json",
                {
                    "rubric_id": interaction.INITIAL_RUBRIC["rubric_id"],
                    "question": "Which qualitative priority should govern this decision?",
                    "interaction_stage": 2,
                },
            )
            (feedback / "human_expert_dialogue.md").write_text(
                "expert dialogue", encoding="utf-8"
            )
            interaction.workflow_evolution.write_json(
                feedback / "expert_reply_1.json",
                {"ok": True, "answer": "Use the safety priority."},
            )
            evidence.write_text('{"stable": true}', encoding="utf-8")
            interaction.workflow_evolution.write_json(
                output / "results/interaction_decision.json",
                {
                    "core_decision": "Choose the governing service priority.",
                    "input_type": "value_preference",
                    "status": "ADOPTED_VALUE",
                    "adopted_scope": "The qualitative priority only.",
                    "comparison_after_two": "stable",
                    "technical_translations": [
                        {
                            "name": "Balanced translation",
                            "evidence_file": "results/comparison.json",
                            "result": "must be discarded",
                        },
                        {
                            "name": "Safety-first translation",
                            "evidence_file": "results/comparison.json",
                        },
                    ],
                    "external_validation": {"required": False, "sources": []},
                },
            )

            receipt = interaction.finalize_interaction_receipt(
                {"run_dir": run_dir}, interaction.INITIAL_RUBRIC
            )

            translations = receipt["adoption_decisions"][0][
                "technical_translations"
            ]
            self.assertEqual(set(translations[0]), {"name", "evidence_file"})
            self.assertEqual(receipt["exchange_count"], 1)
            self.assertEqual(receipt["generated_by"], "python_controller")
            interaction.validate_interaction_run(
                {
                    "run_dir": run_dir,
                    "final_report": self._write_report_after_dialogue(output),
                },
                interaction.INITIAL_RUBRIC,
            )

    @staticmethod
    def _write_report_after_dialogue(output: Path) -> Path:
        report = output / "results/solution_report.md"
        report.write_text("final report", encoding="utf-8")
        return report

    def test_new_run_uses_one_openclaw_execution_with_blocking_expert_bridge(self):
        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary)
            output = run_dir / "output"
            role = output / "logs/operator_roles/human_modeling_expert.md"
            role.parent.mkdir(parents=True, exist_ok=True)
            role.write_text("expert role", encoding="utf-8")
            prepared = {
                "problem_id": interaction.VALIDATION_PROBLEMS[0],
                "repetition": 1,
                "run_dir": run_dir,
                "output_dir": output,
                "prompt": run_dir / "prompt.md",
                "final_report": output / "results/solution_report.md",
                "agent_id": "test-agent",
                "recovered": False,
            }
            args = SimpleNamespace(skip_interaction_receipt_validation=True)

            def fake_openclaw(*_args):
                feedback = output / "logs/operator_feedback"
                interaction.workflow_evolution.write_json(
                    feedback / "expert_request_1.json",
                    {
                        "rubric_id": "Concise task-balanced consultation",
                        "question": "Which qualitative priority should govern?",
                        "interaction_stage": 2,
                    },
                )
                reply = feedback / "expert_reply_1.json"
                deadline = time.monotonic() + 2
                while not reply.is_file() and time.monotonic() < deadline:
                    time.sleep(0.01)
                self.assertTrue(reply.is_file())
                evidence = output / "results/comparison.json"
                evidence.parent.mkdir(parents=True, exist_ok=True)
                evidence.write_text("{}", encoding="utf-8")
                interaction.workflow_evolution.write_json(
                    output / "results/interaction_decision.json",
                    {
                        "core_decision": "Choose one priority.",
                        "input_type": "value_preference",
                        "status": "ADOPTED_VALUE",
                        "adopted_scope": "Qualitative priority only.",
                        "comparison_after_two": "stable",
                        "technical_translations": [
                            {"name": "A", "evidence_file": "results/comparison.json"},
                            {"name": "B", "evidence_file": "results/comparison.json"},
                        ],
                        "external_validation": {"required": False, "sources": []},
                    },
                )
                prepared["final_report"].write_text("report", encoding="utf-8")

            judged = {
                "average_score": 0.8,
                "dimension_scores": {},
                "judge_repeats": 1,
                "judge_stability_result": "stability.json",
                "judge_result": "judge.json",
            }
            with patch.object(
                interaction, "run_local_modeling_phase", side_effect=fake_openclaw
            ) as run_phase, patch.object(
                interaction, "call_direct_human_expert", return_value="Prefer safety."
            ) as expert, patch.object(
                interaction, "judge_report", return_value=judged
            ):
                result = interaction.run_validation_problem(
                    prepared,
                    interaction.INITIAL_RUBRIC,
                    1,
                    run_dir,
                    args,
                )

            run_phase.assert_called_once()
            expert.assert_called_once()
            self.assertEqual(result["interaction_receipt"]["exchange_count"], 1)
            reply = interaction.workflow_evolution.read_json(
                output / "logs/operator_feedback/expert_reply_1.json", {}
            )
            self.assertTrue(reply["rubric_id_normalized"])
            self.assertEqual(
                reply["rubric_id"], interaction.INITIAL_RUBRIC["rubric_id"]
            )

    def test_incomplete_interaction_is_never_scored_even_when_receipt_check_is_skipped(self):
        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary)
            report = run_dir / "output/results/solution_report.md"
            report.parent.mkdir(parents=True, exist_ok=True)
            report.write_text("report without expert reply", encoding="utf-8")
            prepared = {
                "problem_id": interaction.VALIDATION_PROBLEMS[0],
                "repetition": 1,
                "run_dir": run_dir,
                "prompt": run_dir / "prompt.md",
                "final_report": report,
                "recovered": True,
            }
            args = SimpleNamespace(skip_interaction_receipt_validation=True)

            with patch.object(interaction, "judge_report") as judge:
                with self.assertRaisesRegex(
                    RuntimeError, "Expert interaction did not complete"
                ):
                    interaction.run_validation_problem(
                        prepared,
                        interaction.INITIAL_RUBRIC,
                        1,
                        run_dir,
                        args,
                    )

            judge.assert_not_called()

    def test_bridge_validation_error_writes_failure_reply_to_unblock_agent(self):
        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary)
            output = run_dir / "output"
            feedback = output / "logs/operator_feedback"
            role = output / "logs/operator_roles/human_modeling_expert.md"
            role.parent.mkdir(parents=True, exist_ok=True)
            role.write_text("expert role", encoding="utf-8")
            interaction.workflow_evolution.write_json(
                feedback / "expert_request_1.json",
                {
                    "rubric_id": "any-id",
                    "question": "Which value should govern?",
                    "interaction_stage": 99,
                },
            )
            stop = interaction.threading.Event()
            errors = []
            interaction.serve_expert_requests(
                {"output_dir": output, "final_report": output / "results/report.md"},
                interaction.INITIAL_RUBRIC,
                SimpleNamespace(),
                stop,
                errors,
            )

            reply = interaction.workflow_evolution.read_json(
                feedback / "expert_reply_1.json", {}
            )
            self.assertFalse(reply["ok"])
            self.assertIn("interaction_stage", reply["error"])
            self.assertEqual(len(errors), 1)

    def test_evaluate_round_retries_only_failed_problems_serially(self):
        with tempfile.TemporaryDirectory() as temporary:
            experiment = Path(temporary)
            prompt_path = experiment / "prompt.md"
            prompt_path.write_text("prompt", encoding="utf-8")
            attempts = {problem_id: 0 for problem_id in interaction.VALIDATION_PROBLEMS}

            def prepare_problem(
                problem_id,
                _problem,
                _prompt_path,
                _round_number,
                repetition,
                _experiment,
                _args,
            ):
                return {
                    "problem_id": problem_id,
                    "repetition": repetition,
                    "agent_id": None,
                }

            def run_problem(prepared, *_args):
                problem_id = prepared["problem_id"]
                attempts[problem_id] += 1
                if problem_id != interaction.VALIDATION_PROBLEMS[-1] and attempts[problem_id] == 1:
                    raise RuntimeError("temporary gateway timeout")
                return {
                    "problem_id": problem_id,
                    "average_score": 0.75,
                    "dimension_scores": {
                        "structural_coherency": 0.75,
                        "modeling_groundedness": 0.75,
                    },
                }

            args = SimpleNamespace(
                concurrency=3,
                validation_attempts=3,
                retry_concurrency=1,
                validation_retry_delay=0,
            )
            problems = {
                problem_id: {"question": "question"}
                for problem_id in interaction.VALIDATION_PROBLEMS
            }
            with patch.object(
                interaction,
                "prepare_validation_problem",
                side_effect=prepare_problem,
            ), patch.object(
                interaction, "run_validation_problem", side_effect=run_problem
            ):
                result = interaction.evaluate_round(
                    experiment,
                    1,
                    interaction.INITIAL_RUBRIC,
                    problems,
                    prompt_path,
                    args,
                )

            self.assertEqual(result["utility"], 0.75)
            self.assertEqual(
                attempts,
                {
                    interaction.VALIDATION_PROBLEMS[0]: 4,
                    interaction.VALIDATION_PROBLEMS[1]: 4,
                    interaction.VALIDATION_PROBLEMS[2]: 3,
                },
            )
            checkpoint = interaction.workflow_evolution.read_json(
                experiment / "workflows/round_1/evaluation_checkpoint.json", {}
            )
            self.assertEqual(set(checkpoint["completed"]), set(attempts))
            self.assertTrue(
                all(
                    len(problem_runs) == interaction.VALIDATION_REPETITIONS
                    for problem_runs in checkpoint["completed"].values()
                )
            )
            self.assertEqual(checkpoint["failed"], {})

    def test_optimizer_retries_explore_distinct_interaction_axes(self):
        instructions = [
            interaction.optimizer_attempt_instruction(attempt, None)
            for attempt in range(1, 6)
        ]

        self.assertEqual(len(set(instructions)), 5)
        self.assertTrue(
            any("stakeholder" in item and "decision" in item for item in instructions)
        )
        retry = interaction.optimizer_attempt_instruction(
            2, "optimizer returned a previously evaluated rubric"
        )
        self.assertIn("preceding attempt was rejected", retry)
        self.assertIn("changing only the\nname", retry)

    def test_mutation_axis_persists_across_successful_rounds(self):
        results = []
        selected = []
        for round_number in range(2, 7):
            axis, _ = interaction.mutation_axis_for_attempt(results, 1)
            selected.append(axis)
            results.append(
                {
                    "round": round_number,
                    "rubric": {"mutation_axis": axis},
                }
            )

        self.assertEqual(selected, [axis for axis, _ in interaction.MUTATION_AXES])

    def test_mutation_axis_resume_prioritizes_underexplored_axis(self):
        timing_axis = interaction.MUTATION_AXES[0][0]
        results = [
            {"round": round_number, "rubric": {"mutation_axis": timing_axis}}
            for round_number in range(2, 6)
        ]

        first_axis, _ = interaction.mutation_axis_for_attempt(results, 1)
        retry_axis, _ = interaction.mutation_axis_for_attempt(results, 2)

        self.assertEqual(first_axis, interaction.MUTATION_AXES[1][0])
        self.assertEqual(retry_axis, interaction.MUTATION_AXES[2][0])
        self.assertNotEqual(first_axis, timing_axis)

    def test_crossover_parents_prefer_two_evolved_distinct_axes(self):
        def evolved_rubric(name, axis):
            return {
                "rubric_id": f"interaction_{name}",
                "name": name,
                "criteria": [
                    {
                        "name": f"criterion for {name}",
                        "rule": f"operational rule inherited from {name}",
                        "positive_example": f"positive example for {name}",
                        "negative_example": f"negative example for {name}",
                    }
                ],
                "attention_budget": "Use one focused expert exchange and then stop.",
                "mutation_axis": axis,
            }

        timing_axis = interaction.MUTATION_AXES[0][0]
        context_axis = interaction.MUTATION_AXES[1][0]
        results = [
            {
                "round": 1,
                "utility": 0.99,
                "rubric": interaction.INITIAL_RUBRIC,
            },
            {
                "round": 2,
                "utility": 0.90,
                "rubric": evolved_rubric("best", timing_axis),
            },
            {
                "round": 3,
                "utility": 0.89,
                "rubric": evolved_rubric("same_axis", timing_axis),
            },
            {
                "round": 4,
                "utility": 0.80,
                "rubric": evolved_rubric("complementary", context_axis),
            },
        ]

        primary, secondary = interaction.select_evolution_parents(results)

        self.assertEqual(primary["round"], 2)
        self.assertEqual(secondary["round"], 4)

    def test_crossover_parent_selection_falls_back_until_two_evolved_exist(self):
        primary, secondary = interaction.select_evolution_parents(
            [{"round": 1, "utility": 1.0, "rubric": interaction.INITIAL_RUBRIC}]
        )
        self.assertEqual(primary["round"], 1)
        self.assertIsNone(secondary)

    def test_proposed_child_records_two_parent_crossover_and_mutation(self):
        def parent(round_number, axis):
            return {
                "round": round_number,
                "utility": 1.0 - round_number / 100,
                "rubric": {
                    "rubric_id": f"interaction_parent_{round_number}",
                    "name": f"parent rubric {round_number}",
                    "criteria": [
                        {
                            "name": f"parent criterion {round_number}",
                            "rule": "Use one observable operational interaction rule.",
                            "positive_example": "Ask one focused qualitative decision question.",
                            "negative_example": "Ask the expert to perform technical calculations.",
                        }
                    ],
                    "attention_budget": "Use one focused exchange and stop promptly.",
                    "mutation_axis": axis,
                },
                "problem_results": [],
            }

        primary = parent(2, interaction.MUTATION_AXES[0][0])
        secondary = parent(3, interaction.MUTATION_AXES[1][0])
        results = [primary, secondary]
        problems = {
            problem_id: {"title": problem_id, "question": f"Question {problem_id}"}
            for problem_id in interaction.VALIDATION_PROBLEMS
        }
        candidate = {
            "name": "crossed compact consultation",
            "purpose": "Combine complementary parent mechanisms in one consultation.",
            "criteria": [
                {
                    "name": "crossed operational rule",
                    "rule": "Apply both inherited mechanisms in one observable exchange.",
                    "positive_example": "Ask one qualitative question with a clear adoption gate.",
                    "negative_example": "Concatenate unrelated questions from both parents.",
                }
            ],
            "attention_budget": "Use one compact exchange and stop after adoption.",
            "success_test": "The dialogue demonstrates both inherited mechanisms and validation.",
            "mutation_axis": "optimizer supplied value is overwritten",
            "mutation_rationale": "Mutate the assigned interaction mechanism after crossover.",
            "crossover_rationale": (
                "Inherited the question trigger from parent 2 and the adoption gate "
                "from parent 3."
            ),
        }
        args = SimpleNamespace(optimizer_retries=1, opt_model="test-model")

        with patch.object(
            interaction.local, "optimizer_response", return_value=candidate
        ) as optimizer:
            child = interaction.propose_rubric(
                primary, secondary, results, 4, args, problems
            )

        prompt = optimizer.call_args.args[0]["messages"][1]["content"]
        self.assertIn("explicit two-parent crossover", prompt)
        self.assertEqual(child["evolution_operator"], "crossover_mutation")
        self.assertEqual(child["parent_round"], 2)
        self.assertEqual(child["secondary_parent_round"], 3)
        self.assertEqual(child["parent_rounds"], [2, 3])

    def test_optimizer_problem_context_exposes_question_not_judge_configuration(self):
        problems = {
            problem_id: {
                "title": f"Title for {problem_id}",
                "question": f"Complete statement for {problem_id}",
                "eval_roles": ["must remain hidden"],
                "decomposition": {"hidden": True},
            }
            for problem_id in interaction.VALIDATION_PROBLEMS
        }

        context = interaction.optimizer_problem_context(problems)

        self.assertEqual(len(context), 3)
        self.assertTrue(all("question" in item for item in context))
        self.assertTrue(all("eval_roles" not in item for item in context))
        self.assertTrue(all("decomposition" not in item for item in context))

    def test_judge_feedback_summary_uses_only_low_scoring_subdimensions(self):
        with tempfile.TemporaryDirectory() as temporary:
            judge_path = Path(temporary) / "judge.json"
            interaction.workflow_evolution.write_json(
                judge_path,
                {
                    "judgements": {
                        "structural_coherency": {
                            "calculated_overall": 1.0,
                            "scores": {"complete section": 1.0},
                            "explanation": {
                                "complete section": "perfect feedback must stay hidden"
                            },
                        },
                        "scoring_decomposition": {
                            "calculated_overall": 0.75,
                            "scores": {
                                "weak communication": 0.75,
                                "complete model": 1.0,
                            },
                            "explanation": {
                                "weak communication": (
                                    "The presentation has strengths, but its claims are "
                                    "not consistently aligned with reported results."
                                ),
                                "complete model": "high-score detail must stay hidden",
                            },
                        },
                        "analysis_groundedness": {
                            "aggregated_score": 0.5,
                            "role_based_results": [
                                {
                                    "mathematical_rigor": {
                                        "score": 0.5,
                                        "explanation": "Formal justification is incomplete.",
                                    },
                                    "critical_analysis": {
                                        "score": 1.0,
                                        "explanation": "full-score role detail must stay hidden",
                                    },
                                }
                            ],
                        },
                    }
                },
            )
            response = {
                "low_score_dimensions": [
                    {
                        "dimension": "scoring_decomposition",
                        "insufficient_aspects": ["claim consistency"],
                    },
                    {
                        "dimension": "analysis_groundedness",
                        "insufficient_aspects": ["formal justification"],
                    },
                ]
            }
            args = SimpleNamespace(
                opt_model="test-model",
                judge_summary_model=None,
                opt_api_key=None,
                opt_base_url=None,
            )

            with patch.object(
                interaction.local, "optimizer_response", return_value=response
            ) as summarize:
                summary = interaction.summarize_judge_feedback_file(
                    judge_path, args
                )

            prompt = summarize.call_args.args[0]["messages"][1]["content"]
            self.assertEqual(summary, response)
            self.assertIn("not consistently aligned", prompt)
            self.assertIn("Formal justification is incomplete", prompt)
            self.assertNotIn("perfect feedback must stay hidden", prompt)
            self.assertNotIn("high-score detail must stay hidden", prompt)
            self.assertNotIn("full-score role detail must stay hidden", prompt)
            self.assertIn("Output weaknesses only", prompt)

    def test_optimizer_evidence_includes_only_abstract_judge_summary(self):
        problems = {
            problem_id: {
                "title": f"Title for {problem_id}",
                "question": f"Statement for {problem_id}",
            }
            for problem_id in interaction.VALIDATION_PROBLEMS
        }
        results = [
            {
                "round": 2,
                "utility": 0.7,
                "rubric": {"name": "rubric"},
                "problem_results": [
                    {"problem_id": problem_id, "run_dir": "."}
                    for problem_id in interaction.VALIDATION_PROBLEMS
                ],
            }
        ]
        problem_id = interaction.VALIDATION_PROBLEMS[0]
        summaries = {
            (2, problem_id): {
                "low_score_dimensions": [
                    {
                        "dimension": "modeling_groundedness",
                        "insufficient_aspects": ["uncertainty treatment"],
                    }
                ]
            }
        }

        evidence = interaction.optimizer_top_node_evidence(
            results,
            problems,
            judge_weakness_summaries=summaries,
        )

        run = evidence[0]["validation_runs"][0]
        self.assertEqual(run["judge_weakness_summary"], summaries[(2, problem_id)])
        self.assertNotIn("raw_judge_feedback", run)

    def test_optimizer_parent_trajectory_orders_workflow_and_interaction_files(self):
        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary) / "run"
            output = run_dir / "output"
            step = output / "results" / "step_01_problem_understanding.md"
            dialogue = (
                output / "logs" / "operator_feedback" / "human_expert_dialogue.md"
            )
            report = output / "results" / "solution_report.md"
            for path in (step, dialogue, report):
                path.parent.mkdir(parents=True, exist_ok=True)
            step.write_text("step content", encoding="utf-8")
            dialogue.write_text("expert interaction content", encoding="utf-8")
            report.write_text("complete final report", encoding="utf-8")
            for timestamp, path in enumerate((step, dialogue, report), 100):
                os.utime(path, (timestamp, timestamp))
            parent = {
                "round": 1,
                "rubric_id": "rubric_test",
                "problem_results": [
                    {
                        "problem_id": "problem",
                        "run_dir": str(run_dir),
                    }
                ],
            }

            trajectories = interaction.optimizer_parent_trajectories(parent)

            self.assertEqual(
                trajectories,
                [
                    {
                        "problem_id": "problem",
                        "trajectory": [
                            "step content",
                            "expert interaction content",
                            "complete final report",
                        ],
                    }
                ],
            )

    def test_optimizer_uses_top_three_nodes_with_problem_statements_and_trajectories(self):
        with tempfile.TemporaryDirectory() as temporary:
            problems = {
                problem_id: {
                    "title": f"Title for {problem_id}",
                    "question": f"Statement for {problem_id}",
                }
                for problem_id in interaction.VALIDATION_PROBLEMS
            }
            results = []
            for round_number, utility in ((1, 0.4), (2, 0.9), (3, 0.7), (4, 0.8)):
                problem_results = []
                for problem_id in interaction.VALIDATION_PROBLEMS:
                    run_dir = (
                        Path(temporary)
                        / f"round_{round_number}"
                        / baseline.safe_path_component(problem_id)
                    )
                    report = run_dir / "output" / "results" / "solution_report.md"
                    report.parent.mkdir(parents=True, exist_ok=True)
                    report.write_text(
                        f"trajectory round {round_number} for {problem_id}",
                        encoding="utf-8",
                    )
                    problem_results.append(
                        {"problem_id": problem_id, "run_dir": str(run_dir)}
                    )
                results.append(
                    {
                        "round": round_number,
                        "utility": utility,
                        "rubric": {"name": f"rubric {round_number}"},
                        "problem_results": problem_results,
                    }
                )

            evidence = interaction.optimizer_top_node_evidence(results, problems)

            self.assertEqual([item["round"] for item in evidence], [2, 4, 3])
            self.assertEqual([item["rank"] for item in evidence], [1, 2, 3])
            self.assertEqual(
                len(interaction.optimizer_top_node_evidence(results[:2], problems)),
                2,
            )
            self.assertTrue(
                all(len(item["validation_runs"]) == 3 for item in evidence)
            )
            self.assertEqual(
                evidence[0]["validation_runs"][0]["question"],
                f"Statement for {interaction.VALIDATION_PROBLEMS[0]}",
            )
            self.assertIn(
                "trajectory round 2",
                evidence[0]["validation_runs"][0]["trajectory"][0],
            )

    def test_objective_excludes_innovativeness_and_data_groundedness(self):
        dimensions = {
            "structural_coherency": 0.8,
            "modeling_groundedness": 0.6,
            "innovativeness": 0.0,
            "data_groundedness": 0.1,
        }

        self.assertEqual(
            interaction.objective_dimensions(dimensions),
            {"structural_coherency": 0.8, "modeling_groundedness": 0.6},
        )
        self.assertAlmostEqual(interaction.objective_score(dimensions), 0.7)

    def test_legacy_results_are_normalized_before_evolution(self):
        legacy = [
            {
                "round": 1,
                "parent_round": None,
                "rubric": interaction.INITIAL_RUBRIC,
                "utility": 0.5,
                "average_dimension_scores": {},
                "problem_results": [
                    {
                        "problem_id": "problem",
                        "average_score": 0.5,
                        "dimension_scores": {
                            "modeling_groundedness": 0.8,
                            "analysis_groundedness": 0.6,
                            "innovativeness": 0.0,
                            "data_groundedness": 0.0,
                        },
                    }
                ],
            }
        ]

        normalized = interaction.normalize_results(legacy)

        self.assertAlmostEqual(normalized[0]["utility"], 0.7)
        self.assertEqual(
            normalized[0]["problem_results"][0]["dimension_scores"],
            {"modeling_groundedness": 0.8, "analysis_groundedness": 0.6},
        )
        history = interaction.interaction_history(normalized)
        self.assertNotIn(
            "innovativeness", history[0]["per_problem"][0]["dimension_scores"]
        )


if __name__ == "__main__":
    unittest.main()
