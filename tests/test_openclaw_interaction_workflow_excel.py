import json
import os
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from src.OpenClaw import interaction_workflow_excel as archive


SEED_ID = "interaction_workflow_seed0000000"
CANDIDATE_ID = "interaction_workflow_candidate000"


def moment(year, month, day, hour, minute, second) -> float:
    """Local epoch seconds for the naive timestamps the engine writes."""
    return datetime(year, month, day, hour, minute, second).timestamp()


# 2026-09-23 22:54:54 -> 2026-09-24 00:34:52, the round interval the fixtures
# encode: a parent phase, a policy rewrite, a candidate phase.
PARENT_START = moment(2026, 9, 23, 22, 54, 54)
PARENT_DONE = moment(2026, 9, 23, 23, 59, 35)
EVOLUTION_DONE = moment(2026, 9, 23, 23, 59, 59)
CANDIDATE_START = moment(2026, 9, 23, 23, 59, 59)
CANDIDATE_DONE = moment(2026, 9, 24, 0, 34, 52)


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def write_run(root: Path, phase: str, run_dir: Path, created_at: str, *,
              with_metadata: bool = True) -> None:
    """One run directory: metadata (or not, for a stray partial launch)."""
    if with_metadata:
        write_json(
            run_dir / "meta" / "run.json",
            {
                "problem_id": "2017_D",
                "repetition": 1,
                "created_at": created_at,
                "round": 1,
            },
        )
    write_json(
        run_dir / "output" / "results" / "interaction_receipt.json",
        {
            "rubric_id": "interaction_initial_substantive_v1",
            "questions_asked": ["# Expert Question 1 — Which framework?"],
            "expert_answers": ["Endorse the discrete-event simulation."],
            "exchange_count": 1,
            "workflow_id": CANDIDATE_ID,
            "workflow_max_exchanges": 3,
        },
    )
    (run_dir / "output" / "logs" / "operator_feedback").mkdir(
        parents=True, exist_ok=True
    )
    (run_dir / "output" / "logs" / "operator_feedback" / "human_expert_dialogue.md").write_text(
        "# Human Expert Dialogue\n\nQ1 ...\nA1 ...\n", encoding="utf-8"
    )


def build_experiment(root: Path) -> None:
    workflows = root / "workflows"
    workflows.mkdir(parents=True, exist_ok=True)

    write_json(
        root / "initial_interaction_workflows.json",
        [
            {
                "workflow_id": SEED_ID,
                "name": "Strategic decision consultation policy",
                "purpose": "Consult only for high-impact strategic decisions.",
                "policy_text": "# Human Expert Interaction\n\n## Step 1\nAsk once.\n",
                "max_exchanges": 3,
                "stop_condition": "Stop when the direction is settled.",
            }
        ],
    )
    candidate_workflow = {
        "workflow_id": CANDIDATE_ID,
        "name": "Constraint-Aware Strategic Consultation",
        "purpose": "Make the expert name the constraints the agent missed.",
        "policy_text": (
            "# Human Expert Interaction\n\n## Step 1\nAsk once.\n"
            "### Constraint Check\nName the constraints you believe are missing.\n"
        ),
        "max_exchanges": 3,
        "stop_condition": "Stop when the direction is settled.",
        "evolution_mode": "mutation",
        "changed_components": ["Added a 'Constraint Check' section to Step 2."],
        "evolution_rationale": "The expert volunteered constraints unprompted.",
        "actions": [
            {
                "action_id": "step_1_identify_decision_point",
                "action_type": "agent_step",
                "rule": "Summarize the decision point.",
            }
        ],
    }
    round_result = {
        "round": 1,
        "workflow_id": CANDIDATE_ID,
        "workflow": candidate_workflow,
        "completed_at": "2026-09-24T00:34:51.000000",
        "training_batch": ["2017_D", "2017_A"],
        "training_parent_workflow_ids": [SEED_ID],
        "train_pre_utility": 0.8648,
        "train_post_utility": 0.8627,
        "train_accepted": False,
        "validation_utility": None,
        "validation_accepted": None,
        "problem_results": [],
    }
    write_json(workflows / "round_1" / "result.json", round_result)
    write_json(workflows / "round_1" / "workflow.json", candidate_workflow)
    write_json(workflows / "results.json", [round_result])
    write_json(
        workflows / "best_workflow.json",
        {"workflow_id": SEED_ID, "source_round": 0, "validation_utility": 0.8565},
    )
    write_json(
        workflows / "round_1" / "training_parent_evidence.json", {"round": 1}
    )
    write_json(workflows / "round_1" / "evolution_response.json", {"round": 1})
    os.utime(workflows / "round_1" / "training_parent_evidence.json", (PARENT_DONE, PARENT_DONE))
    os.utime(workflows / "round_1" / "evolution_response.json", (EVOLUTION_DONE, EVOLUTION_DONE))
    os.utime(workflows / "round_1" / "result.json", (CANDIDATE_DONE, CANDIDATE_DONE))

    # The seed's own blind-split rollout: the bar the first champion has to clear.
    write_json(
        root / "cpe_evaluations" / "round_0" / archive.SEED_PHASE
        / "workflows" / "round_0" / "result.json",
        {"round": 0, "workflow_id": SEED_ID, "utility": 0.8565},
    )

    parent_run = (
        root / "cpe_evaluations" / "round_1" / "train_parent_1" / "runs" / "round_1"
        / "r1p1_20260923_225454"
    )
    write_run(root, "train_parent_1", parent_run, "2026-09-23T22:54:54.000000")
    os.utime(parent_run / "meta" / "run.json", (PARENT_START, PARENT_START))

    candidate_phase = root / "cpe_evaluations" / "round_1" / "train_candidate"
    candidate_run = candidate_phase / "runs" / "round_1" / "r1p1_20260923_235959"
    write_run(root, "train_candidate", candidate_run, "2026-09-23T23:59:59.000000")
    os.utime(candidate_run / "meta" / "run.json", (CANDIDATE_START, CANDIDATE_START))
    write_json(
        candidate_phase / "workflows" / "round_1" / "result.json",
        {
            "round": 1,
            "workflow_id": CANDIDATE_ID,
            "problem_results": [
                {
                    "problem_id": "2017_D",
                    "average_score": 0.9,
                    "net_utility": 0.8906,
                    "interaction_cost_runs": [
                        {"total_tokens_approx": 3873, "total_expert_latency_seconds": 4.31}
                    ],
                    "repetitions": [
                        {
                            "repetition": 1,
                            "run_dir": str(candidate_run),
                            "average_score": 0.9,
                            "final_report": str(
                                candidate_run / "output" / "results" / "solution_report.md"
                            ),
                            "interaction_receipt": {
                                "questions_asked": ["# Expert Question 1 — Which framework?"],
                                "expert_answers": ["Endorse the discrete-event simulation."],
                                "exchange_count": 1,
                            },
                        }
                    ],
                }
            ],
        },
    )
    # A stray partial directory: an interrupted launch leaves run data but no
    # run metadata, and must not enter the sheet as a blank row.
    (candidate_phase / "runs" / "round_1" / "r1p9_20260923_235959" / "output" / "code").mkdir(
        parents=True, exist_ok=True
    )


class InteractionWorkflowExcelTests(unittest.TestCase):
    def export(self, temporary: str, results=None) -> Path:
        experiment = Path(temporary)
        build_experiment(experiment)
        output = experiment / "workflows" / archive.WORKBOOK_NAME
        archive.write_interaction_evolution_workbook(output, experiment, results)
        return output

    def test_workbook_carries_strategy_history_and_timing_sheets(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary:
            output = self.export(temporary)
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                [archive.SHEET_STRATEGY, archive.SHEET_HISTORY, archive.SHEET_TIMING],
            )

            strategy = workbook[archive.SHEET_STRATEGY]
            self.assertEqual(strategy.max_row, 3)
            self.assertEqual(strategy["A2"].value, archive.SEED_LABEL)
            self.assertEqual(strategy["C2"].value, SEED_ID)
            self.assertEqual(strategy["O2"].value, 0.8565)
            self.assertEqual(strategy["A3"].value, 1)
            self.assertEqual(strategy["C3"].value, CANDIDATE_ID)
            self.assertEqual(strategy["D3"].value, SEED_ID)
            self.assertEqual(strategy["E3"].value, "0")
            self.assertEqual(strategy["N3"].value, "否")
            self.assertAlmostEqual(strategy["M3"].value, 0.8627 - 0.8648, places=6)
            self.assertIn("Constraint Check", strategy["U3"].value)
            self.assertIn("step_1_identify_decision_point", strategy["V3"].value)
            self.assertIn("+ ### Constraint Check", strategy["H3"].value)

            history = workbook[archive.SHEET_HISTORY]
            self.assertEqual(history.max_row, 2)
            self.assertEqual(history["D2"].value, "2017_D")
            self.assertEqual(history["F2"].value, 1)
            self.assertTrue(history["I2"].value.startswith("Q1: # Expert Question 1"))
            self.assertTrue(history["J2"].value.startswith("A1: Endorse"))
            self.assertEqual(history["L2"].value, 3873)
            self.assertAlmostEqual(history["N2"].value, 0.9)
            self.assertIn("Human Expert Dialogue", history["K2"].value)

    def test_stray_partial_run_directory_is_not_exported(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary:
            output = self.export(temporary)
            workbook = load_workbook(output, data_only=True)
            history = workbook[archive.SHEET_HISTORY]
            self.assertEqual(history.max_row, 2)
            directories = [row[15] for row in history.iter_rows(min_row=2, values_only=True)]
            self.assertEqual(len(directories), 1)
            self.assertNotIn("r1p9_", directories[0])

    def test_timing_sheet_splits_the_round_by_phase(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary:
            output = self.export(temporary)
            workbook = load_workbook(output, data_only=True)
            timing = workbook[archive.SHEET_TIMING]
            headers = [cell.value for cell in timing[1]]
            self.assertEqual(headers[:5], ["轮次", "策略名", "轮次开始", "轮次结束", "轮次总耗时(分钟)"])
            row = next(
                values
                for values in timing.iter_rows(min_row=2, values_only=True)
                if values[0] == 1
            )
            self.assertEqual(row[2], "2026-09-23 22:54:54")
            self.assertAlmostEqual(row[5], 64.68, places=1)  # parent phase
            self.assertAlmostEqual(row[6], 0.4, places=1)  # policy rewrite
            self.assertAlmostEqual(row[7], 34.88, places=1)  # candidate phase
            self.assertEqual(row[9], "完成")

    def test_missing_artifacts_still_produce_a_valid_workbook(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary:
            experiment = Path(temporary)
            output = experiment / "workflows" / archive.WORKBOOK_NAME
            archive.write_interaction_evolution_workbook(output, experiment, [])
            workbook = load_workbook(output, data_only=True)

        self.assertEqual(
            workbook.sheetnames,
            [archive.SHEET_STRATEGY, archive.SHEET_HISTORY, archive.SHEET_TIMING],
        )
        self.assertEqual(workbook[archive.SHEET_STRATEGY].max_row, 1)
        self.assertEqual(workbook[archive.SHEET_HISTORY].max_row, 1)

    def test_oversized_question_is_truncated_to_the_cell_limit(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary:
            experiment = Path(temporary)
            build_experiment(experiment)
            path = (
                experiment / "cpe_evaluations" / "round_1" / "train_candidate"
                / "workflows" / "round_1" / "result.json"
            )
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["problem_results"][0]["repetitions"][0]["interaction_receipt"][
                "questions_asked"
            ] = ["Q" + "x" * 40000]
            write_json(path, payload)
            output = experiment / "workflows" / archive.WORKBOOK_NAME
            archive.write_interaction_evolution_workbook(output, experiment, [])
            workbook = load_workbook(output, data_only=True)
            cell = workbook[archive.SHEET_HISTORY]["I2"].value

        self.assertLess(len(cell), 32800)
        self.assertIn("截断", cell)

    def test_each_repetition_reports_its_own_cost(self):
        """A problem-level cost sum used to be shown on every repetition row.

        ``interaction_cost_runs`` holds one entry per repetition, so summing it
        made each row claim its problem's total and any column total double-count.
        """
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary:
            experiment = Path(temporary)
            build_experiment(experiment)
            path = (
                experiment / "cpe_evaluations" / "round_1" / "train_candidate"
                / "workflows" / "round_1" / "result.json"
            )
            payload = json.loads(path.read_text(encoding="utf-8"))
            problem = payload["problem_results"][0]
            problem["interaction_cost_runs"] = [
                {"total_tokens_approx": 2788, "total_expert_latency_seconds": 2.5},
                {"total_tokens_approx": 1218, "total_expert_latency_seconds": 1.5},
            ]
            # A second repetition, built as a copy so the lookup has to fall back
            # to matching on the repetition number rather than on identity.
            problem["repetitions"] = [
                problem["repetitions"][0],
                {
                    **problem["repetitions"][0],
                    "repetition": 2,
                    "average_score": 0.8,
                },
            ]
            write_json(path, payload)
            output = experiment / "workflows" / archive.WORKBOOK_NAME
            archive.write_interaction_evolution_workbook(output, experiment, [])
            workbook = load_workbook(output, data_only=True)
            history = workbook[archive.SHEET_HISTORY]
            tokens = [history["L2"].value, history["L3"].value]
            latency = [history["M2"].value, history["M3"].value]

        self.assertEqual(tokens, [2788, 1218])
        self.assertEqual(latency, [2.5, 1.5])
        self.assertEqual(sum(tokens), 4006, "the column must not double-count")


if __name__ == "__main__":
    unittest.main()
